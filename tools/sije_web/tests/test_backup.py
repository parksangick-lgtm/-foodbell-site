"""하루 한 번 백업 엑셀 시험 — 시트 '행사'(한 줄 = 행사) · '품목'(한 줄 = 품목)."""
import datetime

import openpyxl

from conftest import fake_event
from tools.sije_web.backup import doc_text, write_backup


def test_backup_has_event_and_item_sheets(tmp_path):
    a = fake_event(updatedAt=datetime.datetime(2026, 10, 16, 1, tzinfo=datetime.timezone.utc))
    b = fake_event(place='가짜동 2-2', contacts=['010-0000-0001', '010-0000-0002'])
    out = tmp_path / '백업' / '2026-10-16 시제행사 백업.xlsx'
    write_backup([('a', a), ('b', b)], out)
    wb = openpyxl.load_workbook(out)
    events, items = wb['행사'], wb['품목']
    assert events.max_row == 3                                     # 머리줄 + 2
    head = [c.value for c in events[1]]
    assert head[:3] == ['행사 id', '상태', '행사일']
    row_b = {h: c.value for h, c in zip(head, events[3])}
    assert row_b['장소'] == '가짜동 2-2' and row_b['연락처'] == '010-0000-0001\n010-0000-0002'
    assert items.max_row == 1 + len(a['items']) + len(b['items'])
    assert [c.value for c in items[1]] == ['행사 id', '줄 이름', '구분', '번호', '품목', '수량', '비고']
    assert [c.value for c in items[2]] == ['a', 'L01', '과일', '1', '배', '5개', None]


# ---------- 2단계: 서류 5종 ----------

UTC = datetime.timezone.utc
FAKE_DOCS = [
    ('d1', {'kind': 'sije-quote', 'label': '○○ 김씨 문중 종친회', 'issue': '2026년 10월 17일', 'total': '2,420,000',
            'html': '<div class="brief-row"><dt>수신</dt><dd><span data-edit="">○○ 김씨 문중 종친회</span> 총무님 귀중</dd></div>',
            'createdAt': datetime.datetime(2026, 10, 17, 1, tzinfo=UTC),
            'updatedAt': datetime.datetime(2026, 10, 17, 2, tzinfo=UTC)}),
    ('d2', {'kind': 'staff-pay', 'label': '2026년 10월 3일', 'issue': '2026년 10월 3일', 'total': '150,000',
            'count': '1명', 'html': '<p>가짜</p>',
            'createdAt': datetime.datetime(2026, 10, 3, 1, tzinfo=UTC),
            'updatedAt': datetime.datetime(2026, 10, 3, 1, tzinfo=UTC)}),
]


def test_backup_has_docs_sheet(tmp_path):
    out = tmp_path / '백업.xlsx'
    write_backup([('a', fake_event())], out, FAKE_DOCS)
    ws = openpyxl.load_workbook(out)['서류']
    assert [c.value for c in ws[1]] == ['서류 id', '종류', '구분값', '발행일', '합계', '인원', '만든 시각', '고친 시각', '서류 글자']
    assert ws.max_row == 3
    first = [c.value for c in ws[2]]
    assert first[:8] == ['d1', '시제 견적서', '○○ 김씨 문중 종친회', '2026년 10월 17일', '2,420,000', None,
                         '2026-10-17T01:00:00+00:00', '2026-10-17T02:00:00+00:00']
    assert first[8] == '수신\n○○ 김씨 문중 종친회 총무님 귀중'
    second = [c.value for c in ws[3]]
    assert second[1] == '알바비 지급내역서' and second[5] == '1명'


def test_backup_names_daily_schedule_kind(tmp_path):
    out = tmp_path / '백업.xlsx'
    daily = ('d9', {'kind': 'daily-schedule', 'label': '2026년 10월 21일 (수)', 'issue': '2026년 10월 21일 (수)',
                    'total': '', 'count': '3건', 'html': '<p>가짜동</p>'})
    write_backup([('a', fake_event())], out, [daily])
    row = [c.value for c in openpyxl.load_workbook(out)['서류'][2]]
    assert row[:3] == ['d9', '일일 일정표', '2026년 10월 21일 (수)'] and row[5] == '3건'


def test_backup_without_docs_still_has_empty_docs_sheet(tmp_path):
    out = tmp_path / '백업.xlsx'
    write_backup([('a', fake_event())], out)
    wb = openpyxl.load_workbook(out)
    assert wb['서류'].max_row == 1 and wb['행사'].max_row == 2


def test_backup_survives_control_characters_and_huge_text(tmp_path):
    """붙여넣은 글에 엑셀이 못 받는 글자(제어 문자)가 있어도 백업 전체가 멈추면 안 된다 (검토 M4)."""
    out = tmp_path / '백업.xlsx'
    odd = ('d3', {'kind': 'statement', 'label': '○○\x0b문중', 'issue': '', 'total': '',
                  'html': '<p>가짜\x00동 \x1b문중</p>' + '<p>가</p>' * 40_000})
    write_backup([('a', fake_event())], out, [odd])
    wb = openpyxl.load_workbook(out)
    row = [c.value for c in wb['서류'][2]]
    assert row[2] == '○○문중' and row[8].startswith('가짜동 문중')
    assert len(row[8]) <= 32_767                                # 엑셀 칸 하나의 한도
    assert wb['행사'].max_row == 2                              # 행사 백업도 그대로


def test_doc_text_strips_tags_and_blank_lines():
    html = ('<div>\n  <dt>수신</dt>\n\n  <dd> ○○ 문중 <b>귀중</b> </dd>\n</div><p></p>'
            '<table><tr><td>과실</td><td>대추 &amp; 밤</td></tr></table><p>합계 <span>1,000</span>원</p>')
    assert doc_text(html) == '수신\n○○ 문중 귀중\n과실\n대추 & 밤\n합계 1,000원'
