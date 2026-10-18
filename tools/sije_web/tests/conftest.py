"""시험 공용 준비물 — 시험 프로젝트(foodbell-sije-test)만 쓴다. 손님 정보는 넣지 않는다(전부 가짜)."""
import io
import re
from pathlib import Path

import pytest
import requests
from firebase_admin import auth

from tools.sije_web import allow, fsdb

ROOT = Path(__file__).resolve().parents[3]
CONFIG_JS = ROOT / 'js' / 'sije' / 'firebase-config.js'

# 60만원상 빈 양식의 품목 (손님 정보 아님)
_LEFT = [('과일', '배', '5개', ''), ('과일', '사과', '5개', ''), ('과일', '감', '5개', ''), ('과일', '밤', '5치', ''),
         ('과일', '대추', '5치', ''), ('한과', '약과', '10개', ''), ('한과', '산자', '5장', ''), ('떡', '조청', '1개', ''),
         ('떡', '웃기', '없음', ''), ('떡', '편떡', '3쪽', ''), ('적·생선', '산적', '3쪽', ''), ('적·생선', '계적', '1마리', ''),
         ('적·생선', '조기', '1마리', '큰걸로'), ('적·생선', '두부(소적)', '5쪽', ''), ('전·부침', '녹두부침', '5쪽', ''),
         ('전·부침', '동태전', '10개', ''), ('전·부침', '육원전', '25개', '')]
_RIGHT = [('탕', '북어탕', '1개', '북어,무'), ('탕', '소탕', '1개', '두부,무'), ('탕', '육탕', '1개', '쇠고기,무'),
          ('나물', '도라지', '1개', ''), ('나물', '숙주', '1개', ''), ('나물', '고사리', '1개', ''),
          ('메·갱', '국수', '1그릇', ''), ('메·갱', '밥·국', '', ''), ('기타', '북어포', '1개', ''),
          ('기타', '식혜', '1개', ''), ('기타', '간장·고명', '1개', ''), ('기타', '물김치', '1개', ''), ('기타', '술', '', '')]


def fake_event(**over) -> dict:
    """가짜 행사 한 건 (저장소에 들어가는 모양 그대로). 손님 정보는 전부 가짜."""
    items, no = {}, 1
    for side, rows in (('L', _LEFT), ('R', _RIGHT)):
        for i, (group, name, qty, note) in enumerate(rows, 1):
            items[f'{side}{i:02d}'] = {'group': group, 'no': str(no), 'name': name, 'qty': qty, 'note': note}
            no += 1
    ev = {
        'status': '확정', 'date': '2026-11-13', 'time': '10시', 'sijeil': '음력 10월 3일',
        'expectedDate': '', 'basis': '', 'confirmedText': '',
        'place': '가짜동 1-1 놀이터 위', 'titlePlace': '가짜동 1-1',
        'setting': '7치 1상', 'sansinje': '없음', 'customer': '가짜 문중', 'contacts': ['010-0000-0000'],
        'staffCar': '가짜 담당 · 1호차', 'meal': '20인분', 'amount': '600,000원', 'amountDetail': '상차림 60만',
        'departTime': '07시', 'arriveTime': '09시 30분',
        'mustDo': '산신제 상은 따로 차린다\n도착하면 먼저 전화',
        'templateType': '60만원상', 'items': items, 'fileLabel': '', 'extraItems': '',
        'memo': '주차는 가짜 마을회관 앞', 'sourceFile': '',
    }
    ev.update(over)
    return ev


def fake_jpeg(w=1200, h=900, color=(176, 141, 87)) -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new('RGB', (w, h), color).save(buf, 'JPEG', quality=90)
    return buf.getvalue()


FAKE_STEM = '2026년11월13일 가짜동 1-1 시제행사 내역서'
FAKE_DATE_TEXT = '2026년 11월 13일 (금) 10시'
FAKE_PLACE = '가짜동 1-1 놀이터 위'


def _build_sheet(dest: Path, *, title_place='가짜동 1-1', date_text=FAKE_DATE_TEXT, contacts=1, amount='600,000원',
                 drop_label=None, rename=None, rename_group=None, extra=None,
                 must=('산신제 상은 따로 차린다', '도착하면 먼저 전화'),
                 memo='주차는 가짜 마을회관 앞', with_image=False, items_title='■ 준비 품목 — 60만원상') -> Path:
    """실제 2026 내역서와 같은 모양의 가짜 엑셀 (jobsheet.py 틀): A·B 이름표 → C 값, G~I 이름표 → J 값,
    연락처가 여럿이면 '연락처 2 · 연락처 3' 줄이 끼어 아래가 한 줄씩 밀린다. 손님 정보는 전부 가짜."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '작업지시서'

    def put(rng, value):
        if ':' in rng:
            ws.merge_cells(rng)
        ws[rng.split(':')[0]] = value

    put('A1:M1', f'{title_place}\n{date_text}')
    put('A3:M3', '■ 행사 정보')
    put('A4:B4', '행사일')
    put('C4:M4', date_text)
    phones = [f'010-0000-000{i}' for i in range(contacts)]
    pairs = [('장소', FAKE_PLACE, '시제일', '음력 10월 3일'), ('상차림', '7치 1상', '산신제', '없음'),
             ('고객', '가짜 문중', '연락처', phones[0])]
    if contacts > 1:
        pairs.append(('연락처 2', phones[1], '연락처 3', phones[2] if contacts > 2 else ''))
    pairs += [('담당·차량', '가짜 담당 · 1호차', '식사', '20인분'), ('금액', amount, '금액 내역', '상차림 60만'),
              ('출발 시간', '07시', '도착 시간', '09시 30분')]
    if rename:   # 손으로 다듬은 파일처럼 이름표만 다르게 (예: 시제일 → 지내는 법)
        pairs = [tuple(rename[1] if x == rename[0] else x for x in p) for p in pairs]
    r = 5
    for a, b, c, d in pairs:
        if a != drop_label:
            put(f'A{r}:B{r}', a)
            put(f'C{r}:F{r}', b)
        if c != drop_label:
            put(f'G{r}:I{r}', c)
            put(f'J{r}:M{r}', d)
        r += 1
    r += 1
    put(f'A{r}:M{r}', '■ 꼭 지킬 것')
    for line in must:
        r += 1
        put(f'A{r}:M{r}', line)
    r += 2
    put(f'A{r}:M{r}', items_title)
    r += 1
    for cols in ('ABCDEF', 'HIJKLM'):
        for col, h in zip(cols, ['구분', '번호', '품목', '수량', '비고', '준비']):
            ws[f'{col}{r}'] = h
    top = r + 1
    no = 1
    sides = [('ABCDEF', _LEFT), ('HIJKLM', _RIGHT)]
    if rename_group:   # 예: 전·부침 → '추가 적\n(2번째)' (실제 파일에 있는 묶음 이름)
        sides = [(cols, [((rename_group[1] if g == rename_group[0] else g), *rest) for g, *rest in rows])
                 for cols, rows in sides]
    n = max(len(_LEFT), len(_RIGHT))
    for cols, rows in sides:
        g, b, c, d, e, f = cols
        prev = None
        for i, (group, name, qty, note) in enumerate(rows):
            rr = top + i
            if group != prev:
                ws[f'{g}{rr}'] = group
                prev = group
            ws[f'{b}{rr}'] = no
            ws[f'{c}{rr}'] = name
            ws[f'{d}{rr}'] = qty or None
            ws[f'{e}{rr}'] = note or None
            ws[f'{f}{rr}'] = '☐'
            no += 1
        if len(rows) < n:
            ws[f'{g}{top + len(rows)}'] = '추가\n기입'
            if extra:
                ws[f'{c}{top + len(rows)}'], ws[f'{d}{top + len(rows)}'] = extra
            for rr in range(top + len(rows), top + n):
                ws[f'{f}{rr}'] = '☐'
    r = top + n + 1
    put(f'A{r}:M{r}', '■ 메모')
    memo_row = r + 1
    put(f'A{memo_row}:M{memo_row}', memo)
    put(f'A{memo_row + 1}:M{memo_row + 1}', None)
    put(f'A{memo_row + 3}:M{memo_row + 3}', '원본: 가짜.xls  ·  노란 칸은 손으로 적어 넣는 칸')
    if with_image:
        import random
        from openpyxl.drawing.image import Image as XlImage
        from PIL import Image
        random.seed(2)
        img = Image.frombytes('RGB', (1800, 1350), bytes(random.getrandbits(8) for _ in range(1800 * 1350 * 3)))
        buf = io.BytesIO()
        img.save(buf, 'PNG')
        buf.seek(0)
        ws.add_image(XlImage(buf), f'H{memo_row}')
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest


def make_fake_sheet(tmp_path, *, contacts=1, with_image=False, date_text=FAKE_DATE_TEXT, status='확정',
                    stem=FAKE_STEM, **kw) -> Path:
    """tmp_path/2026년/{status}/{stem}.xlsx 에 가짜 내역서를 만든다."""
    return _build_sheet(tmp_path / '2026년' / status / f'{stem}.xlsx', contacts=contacts, with_image=with_image,
                        date_text=date_text, **kw)


def make_fake_blank(tmp_path, name='60만원상') -> Path:
    """빈 양식 (품목만 쓴다)."""
    return _build_sheet(tmp_path / f'빈 양식 시제행사 내역서 ({name}).xlsx', title_place='시제행사 내역서', date_text='',
                        must=('',), memo='', items_title=f'■ 준비 품목 — {name}')


def make_fake_list(path: Path, rows) -> Path:
    """확정 · 미확정 목록: (날짜, 장소, 전화번호, 2026 예상 날짜, 근거, 확정 날짜·시간) 줄들."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['날짜', '장소', '전화번호', '2026 예상 날짜', '근거 (작년 기준)', '확정 날짜·시간\n(전화 후 적기)'])
    for row in rows:
        ws.append(list(row))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


@pytest.fixture(scope='session')
def repo_url():
    """저장소 폴더를 이 컴퓨터 안에서만 여는 임시 서버 주소."""
    from tools.sije_web.render_pdf import serve_repo
    with serve_repo(ROOT) as url:
        yield url


def _test_project_config() -> dict:
    """firebase-config.js 의 TEST 묶음 값(apiKey · projectId …) — 설정값은 그 파일 한 곳에만 둔다."""
    block = re.search(r'const TEST = \{(.*?)\};', CONFIG_JS.read_text(encoding='utf-8'), re.S).group(1)
    return dict(re.findall(r"(\w+): '([^']*)'", block))


@pytest.fixture(scope='session')
def test_db():
    if not fsdb.KEY_TEST.exists():
        pytest.skip(f'시험 열쇠가 없습니다: {fsdb.KEY_TEST}')
    return fsdb.open_db(fsdb.KEY_TEST)


@pytest.fixture(scope='session')
def custom_token(test_db):
    """custom_token(email, allowed) → 화면의 signInWithCustomToken 에 넣을 토큰. allowed 면 허락 목록에 넣는다."""
    app = fsdb.app_for(fsdb.KEY_TEST)
    used = []

    def make(email: str, allowed: bool) -> str:
        try:
            user = auth.get_user_by_email(email, app=app)
        except auth.UserNotFoundError:
            user = auth.create_user(email=email, email_verified=True, app=app)
        (allow.add if allowed else allow.remove)(test_db, email)
        used.append(email)
        return auth.create_custom_token(user.uid, app=app).decode()

    yield make
    for email in used:
        allow.remove(test_db, email)


@pytest.fixture(scope='session')
def id_token(custom_token):
    """id_token(email, allowed) → 그 메일로 로그인한 사람의 ID 토큰 (REST 로 저장소를 직접 부를 때)."""
    api_key = _test_project_config()['apiKey']

    def make(email: str, allowed: bool) -> str:
        r = requests.post('https://identitytoolkit.googleapis.com/v1/accounts:signInWithCustomToken',
                          params={'key': api_key}, json={'token': custom_token(email, allowed), 'returnSecureToken': True},
                          timeout=30)
        r.raise_for_status()
        return r.json()['idToken']

    return make


@pytest.fixture(scope='module')
def browser():
    """헤드리스 크롬 — 시험 파일(모듈)마다 하나. render_pdf 가 따로 여는 것과 겹치지 않게 모듈 단위로 닫는다."""
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        b = p.chromium.launch()
        yield b
        b.close()


@pytest.fixture
def page(browser):
    """PC 창(1280px) 하나 — 시험마다 새 창이라 로그인 · 저장소 사본이 섞이지 않는다."""
    ctx = browser.new_context(viewport={'width': 1280, 'height': 900})
    pg = ctx.new_page()
    yield pg
    ctx.close()


class FakeCalendar:
    """캘린더 부품(list_tagged · insert · update · delete)의 메모리 가짜 — calls 에 한 일을 쌓는다.
    fail_ids 에 든 일정 번호는 만들기 · 고치기가 실패한다. untagged 는 사장님이 손으로 만든 일정 자리."""

    def __init__(self, fail_ids=frozenset()):
        self.events = {}
        self.untagged = {}
        self.calls = []
        self.fail_ids = set(fail_ids)

    def list_tagged(self):
        return {e['extendedProperties']['private']['foodbellId']: dict(e) for e in self.events.values()}

    def insert(self, body):
        self.calls.append(('insert', body['id']))
        if body['id'] in self.fail_ids:
            raise RuntimeError('흉내: 만들기 실패')
        self.events[body['id']] = dict(body)

    def update(self, event_id, body):
        self.calls.append(('update', event_id))
        if event_id in self.fail_ids:
            raise RuntimeError('흉내: 고치기 실패')
        self.events[event_id] = dict(body)

    def delete(self, event_id):
        self.calls.append(('delete', event_id))
        del self.events[event_id]
