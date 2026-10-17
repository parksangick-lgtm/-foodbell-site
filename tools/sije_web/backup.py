r"""전체 행사 → 엑셀 한 파일 (하루 한 번, '…\시제행사 내역서\백업\YYYY-MM-DD 시제행사 백업.xlsx').
시트 '행사' 한 줄 = 행사 하나, 시트 '품목' 한 줄 = 품목 하나. 그림은 넣지 않는다 (PDF 에 있다)."""
import datetime
from html.parser import HTMLParser
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

CELL_MAX = 32_767   # 엑셀 칸 하나에 들어가는 글자 수 한도

# 설계서 4절 순서
EVENT_COLS = [
    ('status', '상태'), ('date', '행사일'), ('time', '시간'), ('sijeil', '시제일'), ('expectedDate', '예상 날짜'),
    ('basis', '근거'), ('confirmedText', '확정 날짜·시간'), ('place', '장소'), ('titlePlace', '맨 위 주소'),
    ('setting', '상차림'), ('sansinje', '산신제'), ('customer', '고객'), ('contacts', '연락처'),
    ('staffCar', '담당·차량'), ('meal', '식사'), ('amount', '금액'), ('amountDetail', '금액 내역'),
    ('departTime', '출발 시간'), ('arriveTime', '도착 시간'), ('mustDo', '꼭 지킬 것'), ('templateType', '상 종류'),
    ('fileLabel', '파일 이름 글자'), ('extraItems', '추가 기입'), ('memo', '메모'), ('sourceFile', '옮겨 온 엑셀'),
    ('updatedAt', '고친 시각'),
]
ITEM_COLS = ['group', 'no', 'name', 'qty', 'note']

# 2단계: 서류 5종 (js/docs/doc-model.js 의 KINDS 제목과 같게)
DOC_KINDS = {'sije-quote': '시제 견적서', 'event-quote': '행사 견적서', 'sije-worksheet': '시제 작업지시서',
             'statement': '거래명세서', 'staff-pay': '알바비 지급내역서'}
DOC_COLS = [('label', '구분값'), ('issue', '발행일'), ('total', '합계'), ('count', '인원'),
            ('createdAt', '만든 시각'), ('updatedAt', '고친 시각')]
_BLOCK = {'div', 'p', 'br', 'tr', 'td', 'th', 'li', 'dt', 'dd', 'h1', 'h2', 'h3', 'h4', 'table', 'section', 'article'}


class _Text(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag in _BLOCK:
            self.parts.append('\n')

    def handle_endtag(self, tag):
        if tag in _BLOCK:
            self.parts.append('\n')

    def handle_data(self, data):
        self.parts.append(data)


def doc_text(html: str) -> str:
    """서류 전체 모양(html) → 글자만. 칸 · 줄마다 한 줄, 빈 줄 없이."""
    p = _Text()
    p.feed(html or '')
    p.close()
    lines = (' '.join(line.split()) for line in ''.join(p.parts).split('\n'))
    return '\n'.join(line for line in lines if line)


def _clean(s: str) -> str:
    """붙여넣은 글의 제어 문자 하나가 백업 전체를 멈추지 않게 빼고, 엑셀 칸 한도로 자른다."""
    return ILLEGAL_CHARACTERS_RE.sub('', s)[:CELL_MAX]


def _cell(v):
    if isinstance(v, list):
        return _clean('\n'.join(str(x) for x in v))
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    if isinstance(v, str):
        v = _clean(v)
    return v if v not in ('',) else None


def write_backup(events: list[tuple[str, dict]], out: Path, docs: list[tuple[str, dict]] = ()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '행사'
    ws.append(['행사 id', *(h for _, h in EVENT_COLS)])
    items = wb.create_sheet('품목')
    items.append(['행사 id', '줄 이름', '구분', '번호', '품목', '수량', '비고'])
    for id, ev in events:
        ws.append([id, *(_cell(ev.get(k)) for k, _ in EVENT_COLS)])
        for key in sorted(ev.get('items') or {}):
            it = ev['items'][key]
            items.append([id, key, *(_cell(it.get(k)) for k in ITEM_COLS)])
    sheet = wb.create_sheet('서류')
    sheet.append(['서류 id', '종류', *(h for _, h in DOC_COLS), '서류 글자'])
    for id, doc in docs:
        sheet.append([id, DOC_KINDS.get(doc.get('kind'), doc.get('kind')), *(_cell(doc.get(k)) for k, _ in DOC_COLS),
                      _cell(doc_text(doc.get('html', '')))])
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_name(out.stem + '.만드는중.xlsx')
    wb.save(tmp)
    tmp.replace(out)
