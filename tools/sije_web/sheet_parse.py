"""내역서 엑셀 한 장 → 행사 dict (설계서 4절의 칸).

좌표가 아니라 이름표를 찾아 그 오른쪽 합친 칸의 값을 읽는다 — 연락처가 여럿이면 줄이 끼어 아래가 밀리고,
손으로 다듬은 파일도 있어서다 (A·B 열 이름표 → C 열 값, G~I 열 이름표 → J 열 값).
엑셀은 읽기만 한다. 저장하면 붙여 넣은 그림이 사라진다 — load_workbook 뒤 save() 를 부르지 않는다.
"""
import hashlib
import io
import re
import zipfile
from pathlib import Path

import openpyxl
from PIL import Image

STATUSES = ('확정', '미확정', '주문없음')
TEXT_LABELS = {
    '장소': 'place', '시제일': 'sijeil', '상차림': 'setting', '산신제': 'sansinje', '고객': 'customer',
    '담당·차량': 'staffCar', '식사': 'meal', '금액': 'amount', '금액 내역': 'amountDetail',
    '출발 시간': 'departTime', '도착 시간': 'arriveTime',
}
# 사장님 결정(2026-10-16): 이 이름표가 없는 파일(손으로 다듬은 것)은 그 칸을 빈 칸으로 옮긴다
OPTIONAL = {'시제일', '식사', '금액 내역'}
REQUIRED = [x for x in ['행사일', *TEXT_LABELS, '연락처'] if x not in OPTIONAL]
# 사장님 결정(2026-10-16): 시제일 자리에 적힌 다른 이름표는 메모 맨 앞 줄로 ('지내는 법 — …')
TO_MEMO = {'지내는 법', '지내는 곳'}
DATE = re.compile(r'\s*(\d{4})\s*년\s*(\d{1,2})\s*월\s*(\d{1,2})\s*일')


class SheetError(Exception):
    def __init__(self, file: str, reason: str):
        super().__init__(f'{file}: {reason}')
        self.file, self.reason = file, reason


def parse_date_text(text: str) -> tuple[str, str]:
    """'2026년 11월 13일 (금) 10시' → ('2026-11-13', '10시'). 연 · 월 · 일 중 하나라도 없으면 ValueError."""
    m = DATE.match(str(text or ''))
    if not m:
        raise ValueError(f'날짜를 못 읽음: {text!r}')
    y, mo, d = (int(x) for x in m.groups())
    rest = str(text)[m.end():].strip()
    rest = re.sub(r'^\([^)]*\)\s*', '', rest)          # 요일 (금)
    return f'{y:04d}-{mo:02d}-{d:02d}', rest.strip()


def split_filename(stem: str) -> tuple[str, str]:
    """'2026년11월1일  가짜면 60만원추가 시제행사 내역서' → ('2026년11월1일', '  가짜면 60만원추가') — 공백까지 그대로."""
    m = re.match(r'(\d{4}년\d{1,2}월\d{1,2}일)(.*?)(?: 시제행사 내역서)?$', stem)
    if not m:
        raise ValueError(f'파일 이름에서 날짜를 못 찾음: {stem!r}')
    return m.group(1), m.group(2)


def event_id(source_file: str) -> str:
    """파일 이름의 sha1 앞 12자 — 다시 돌려도 같은 id (옮기기를 두 번 해도 행사가 둘이 되지 않게)."""
    return hashlib.sha1(source_file.encode('utf-8')).hexdigest()[:12]


def _text(v) -> str:
    return '' if v is None else str(v)


def _sections(ws) -> dict[str, int]:
    """'■ 행사 정보' 같은 구역 줄 번호. 이름은 ■ 다음 낱말 (행사 · 꼭 · 준비 · 메모)."""
    out = {}
    for r in range(1, ws.max_row + 1):
        v = ws[f'A{r}'].value
        if isinstance(v, str) and v.startswith('■'):
            out.setdefault(v.split()[1] if len(v.split()) > 1 else '', r)
    return out


def parse_items(ws, file: str = '') -> tuple[dict, str, str, int]:
    """품목 → (items {L01: {group, no, name, qty, note}}, 추가 기입 글자, 품목 제목 뒤 글자, 품목표 줄 수)."""
    secs = _sections(ws)
    if '준비' not in secs:
        raise SheetError(file, "'■ 준비 품목' 구역이 없음")
    title = ws[f'A{secs["준비"]}'].value
    template = re.sub(r'^■\s*준비\s*품목', '', title).strip()
    template = re.sub(r'^—\s*', '', template)
    head = next((r for r in range(secs['준비'] + 1, ws.max_row + 1)
                 if ws[f'A{r}'].value == '구분' and ws[f'C{r}'].value == '품목'), None)
    if head is None:
        raise SheetError(file, "품목 머리줄(구분 · 번호 · 품목 …)이 없음")
    end = min([r for r in secs.values() if r > head] or [ws.max_row + 1])
    items, extra = {}, []
    for side, (g, b, c, d, e) in (('L', 'ABCDE'), ('R', 'HIJKL')):
        group, in_extra, n = '', False, 0
        for r in range(head + 1, end):
            gv = _text(ws[f'{g}{r}'].value)
            if gv:
                # '추가 기입' 칸만 손으로 적는 칸이다 — '추가 적 (2번째)' 같은 진짜 묶음과 헷갈리지 않게 글자가 같을 때만
                if re.sub(r'\s+', '', gv) == '추가기입':
                    in_extra = True
                else:
                    group = gv
            cells = [_text(ws[f'{x}{r}'].value) for x in (b, c, d, e)]
            if in_extra:
                words = ' '.join(x for x in cells if x)
                if words:
                    extra.append(words)
            elif any(cells):
                n += 1
                items[f'{side}{n:02d}'] = {'group': group, 'no': cells[0], 'name': cells[1], 'qty': cells[2], 'note': cells[3]}
    rows = sum(1 for r in range(head + 1, end) if any(ws[f'{c}{r}'].value not in (None, '') for c in 'ABCDEFHIJKLM'))
    return items, '\n'.join(extra), template, rows


def parse_sheet(path: Path) -> dict:
    path = Path(path)
    file = path.name
    ws = openpyxl.load_workbook(path, data_only=True).active    # 읽기만
    secs = _sections(ws)
    for need in ('행사', '꼭', '준비', '메모'):
        if need not in secs:
            raise SheetError(file, f"'■ {need}…' 구역이 없음")

    found, contacts, memo_first = {}, [], []
    for r in range(secs['행사'] + 1, secs['꼭']):
        for lab_col, val_col in (('A', 'C'), ('G', 'J')):
            label = ws[f'{lab_col}{r}'].value
            if not isinstance(label, str) or not label.strip():
                continue
            label = label.strip()
            value = _text(ws[f'{val_col}{r}'].value)
            if label == '행사일':
                found['행사일'] = value
            elif label in TEXT_LABELS:
                found[label] = value
            elif re.fullmatch(r'연락처( \d+)?', label):
                found['연락처'] = True
                contacts.append(value)
            elif label in TO_MEMO:
                if value:
                    memo_first.append(f'{label} — {value}')
            else:
                raise SheetError(file, f"모르는 이름표 '{label}'")
    missing = [x for x in REQUIRED if x not in found]
    if missing:
        raise SheetError(file, '이름표 없음: ' + ', '.join(f"'{x}'" for x in missing))
    try:
        date, time = parse_date_text(found['행사일'])
    except ValueError:
        raise SheetError(file, "'행사일' 칸에서 날짜를 못 읽음") from None

    a1 = _text(ws['A1'].value)
    m = DATE.search(a1)
    title_place = (a1[:m.start()] if m else a1).rstrip()

    must = [_text(ws[f'A{r}'].value) for r in range(secs['꼭'] + 1, secs['준비'])]
    items, extra, template, item_rows = parse_items(ws, file)
    memo = list(memo_first)
    for r in range(secs['메모'] + 1, ws.max_row + 1):
        v = _text(ws[f'A{r}'].value)
        if v.startswith('원본:'):
            break
        if v:
            memo.append(v)

    status = path.parent.name
    if status not in STATUSES:
        raise SheetError(file, f"상태 폴더가 아님: '{status}'")
    ev = {k: found.get(label, '') for label, k in TEXT_LABELS.items()}
    ev.update({
        'status': status, 'date': date, 'time': time, 'expectedDate': '', 'basis': '', 'confirmedText': '',
        'titlePlace': title_place, 'contacts': contacts, 'mustDo': '\n'.join(x for x in must if x),
        'templateType': template, 'items': items, 'itemRows': item_rows, 'fileLabel': split_filename(path.stem)[1],
        'extraItems': extra, 'memo': '\n'.join(memo), 'sourceFile': file,
    })
    return ev


def parse_template(path: Path) -> dict:
    """빈 양식 → {name, items}. 이름은 파일 이름 괄호 안 (60만원상 · 7치 고임상 · 1자 고임상)."""
    path = Path(path)
    m = re.search(r'\(([^)]+)\)', path.stem)
    if not m:
        raise SheetError(path.name, '파일 이름에 상 종류(괄호)가 없음')
    items, _, _, rows = parse_items(openpyxl.load_workbook(path, data_only=True).active, path.name)   # 읽기만
    return {'name': m.group(1), 'items': items, 'itemRows': rows}


def extract_images(path: Path) -> list[bytes]:
    """xlsx 안의 xl/media/* 그림을 순서대로 (엑셀 파일은 열기만 한다)."""
    with zipfile.ZipFile(path) as z:
        names = sorted(n for n in z.namelist() if n.startswith('xl/media/'))
        return [z.read(n) for n in names]


def shrink_jpeg(data: bytes, limit: int = 900_000) -> bytes:
    """긴 변 1600px 이하 JPEG 로, limit 바이트 이하가 될 때까지 화질을 0.1 씩 낮춘다 (투명한 곳은 흰색)."""
    im = Image.open(io.BytesIO(data))
    if im.mode in ('RGBA', 'LA', 'P'):
        im = im.convert('RGBA')
        bg = Image.new('RGB', im.size, 'white')
        bg.paste(im, mask=im.split()[-1])
        im = bg
    else:
        im = im.convert('RGB')
    im.thumbnail((1600, 1600))
    for q in range(85, 0, -10):
        buf = io.BytesIO()
        im.save(buf, 'JPEG', quality=q)
        if buf.tell() <= limit:
            return buf.getvalue()
    raise ValueError('그림을 줄여도 너무 큽니다')
